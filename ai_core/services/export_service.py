from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus.tables import Table, TableStyle
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from io import BytesIO, StringIO
import re
import pandas as pd


# =========================================================
# CLEAN MARKDOWN
# =========================================================

def clean_markdown(text: str) -> str:
    """
    Convert markdown-ish response into
    PDF-friendly formatted text.
    """

    # Bold markdown
    text = re.sub(
        r"\*\*(.*?)\*\*",
        r"<b>\1</b>",
        text
    )

    # Bullet cleanup
    text = text.replace("•", "&bull;")

    return text


# =========================================================
# GENERATE MARKDOWN EXPORT
# =========================================================

def generate_markdown(answer: str, sources: list):
    """
    Export AI response as Markdown.

    Why Markdown?
    -------------
    - GitHub compatible
    - Obsidian compatible
    - Lightweight knowledge sharing
    - Developer friendly
    """

    markdown_content = []

    # =====================================================
    # TITLE
    # =====================================================

    markdown_content.append("# AI Response\n")

    # =====================================================
    # MAIN ANSWER
    # =====================================================

    markdown_content.append(answer)

    # =====================================================
    # SOURCES
    # =====================================================

    markdown_content.append("\n\n# Sources\n")

    for i, source in enumerate(sources, start=1):

        markdown_content.append(
            f"## Source {i}"
        )

        markdown_content.append(
            f"- File: {source.get('source', 'Unknown')}"
        )

        markdown_content.append(
            f"- Score: {source.get('score', 0):.3f}"
        )

        markdown_content.append(
            f"- Section: {source.get('section', 'General')}"
        )

        markdown_content.append(
            f"\n{source.get('content', '')[:500]}\n"
        )

    # =====================================================
    # CONVERT TO BUFFER
    # =====================================================

    content = "\n".join(markdown_content)

    buffer = BytesIO()

    buffer.write(content.encode("utf-8"))

    buffer.seek(0)

    return buffer

# =========================================================
# GENERATE EXCEL EXPORT
# =========================================================

def generate_excel(answer: str, sources: list):
    """
    Export AI response + sources to Excel workbook.

    Sheets:
    -------
    1. AI Answer
    2. Retrieval Sources
    """

    output = BytesIO()

    # =====================================================
    # CREATE EXCEL WRITER
    # =====================================================

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        # =================================================
        # SHEET 1 — AI ANSWER
        # =================================================

        answer_df = pd.DataFrame({
            "AI Response": [answer]
        })

        answer_df.to_excel(
            writer,
            sheet_name="AI_Response",
            index=False
        )

        # =================================================
        # SHEET 2 — SOURCES
        # =================================================

        source_rows = []

        for source in sources:

            source_rows.append({

                "File": source.get(
                    "source",
                    "Unknown"
                ),

                "Section": source.get(
                    "section",
                    "General"
                ),

                "Score": source.get(
                    "score",
                    0
                ),

                "Content": source.get(
                    "content",
                    ""
                )
            })

        sources_df = pd.DataFrame(source_rows)

        sources_df.to_excel(
            writer,
            sheet_name="Sources",
            index=False
        )

    # =====================================================
    # FINALIZE BUFFER
    # =====================================================

    output.seek(0)

    return output

# =========================================================
# GENERATE CSV EXPORT
# =========================================================

def generate_csv(answer: str, sources: list):
    """
    Export AI response into CSV format.

    Structure:
    ----------
    Row 1:
        AI Answer

    Remaining rows:
        Sources metadata
    """

    # =====================================================
    # CREATE TABULAR DATA
    # =====================================================

    rows = []

    # Main answer row
    rows.append({
        "Type": "Answer",
        "Content": answer,
        "Source": "",
        "Score": ""
    })

    # =====================================================
    # ADD SOURCES
    # =====================================================

    for src in sources:

        rows.append({
            "Type": "Source",
            "Content": src.get("content", ""),
            "Source": src.get("source", ""),
            "Score": src.get("score", "")
        })

    # =====================================================
    # CREATE DATAFRAME
    # =====================================================

    df = pd.DataFrame(rows)

    # =====================================================
    # SAVE TO MEMORY BUFFER
    # =====================================================

    from io import BytesIO

    buffer = BytesIO()

    df.to_csv(
        buffer,
        index=False
    )

    buffer.seek(0)

    return buffer

# =========================================================
# PDF EXPORT
# =========================================================

def generate_pdf(answer, sources):

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=50,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    story = []

    # =====================================================
    # TITLE
    # =====================================================

    title = Paragraph(
        "<font size=22><b>AI Knowledge Assistant</b></font>",
        styles["Title"]
    )

    story.append(title)
    story.append(Spacer(1, 24))

    # =====================================================
    # ANSWER HEADING
    # =====================================================

    heading = Paragraph(
        "<font size=16><b>Response</b></font>",
        styles["Heading2"]
    )

    story.append(heading)
    story.append(Spacer(1, 12))

    # =====================================================
    # CLEAN ANSWER
    # =====================================================

    cleaned_answer = clean_markdown(answer)

    # Split into paragraphs
    paragraphs = cleaned_answer.split("\n\n")

    for para in paragraphs:

        para = para.strip()

        if not para:
            continue

        p = Paragraph(
            f"<font size=11>{para}</font>",
            styles["BodyText"]
        )

        story.append(p)
        story.append(Spacer(1, 10))

    # =====================================================
    # SOURCES SECTION
    # =====================================================

    if sources:

        story.append(Spacer(1, 20))

        sources_heading = Paragraph(
            "<font size=16><b>Sources</b></font>",
            styles["Heading2"]
        )

        story.append(sources_heading)
        story.append(Spacer(1, 10))

        table_data = [["File", "Chunk", "Score"]]

        for src in sources:

            file_name = (
                src.get("source", "Unknown")
                .split("\\")[-1]
            )

            chunk_id = str(
                src.get("chunk_id", "-")
            )

            score = str(
                round(src.get("score", 0) * 100, 1)
            ) + "%"

            table_data.append([
                file_name,
                chunk_id,
                score
            ])

        table = Table(
            table_data,
            colWidths=[250, 100, 100]
        )

        table.setStyle(TableStyle([

            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),

            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),

            ("GRID", (0, 0), (-1, -1), 1, colors.grey),

            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),

            ("BOTTOMPADDING", (0, 0), (-1, 0), 10),

            ("BACKGROUND", (0, 1), (-1, -1), colors.whitesmoke),
        ]))

        story.append(table)

    # =====================================================
    # BUILD PDF
    # =====================================================

    doc.build(story)

    buffer.seek(0)

    return buffer

# =========================================================
# CLEAN MARKDOWN FOR DOCX
# =========================================================

def clean_docx_markdown(text: str) -> str:
    """
    Prepare markdown-ish AI response
    for DOCX rendering.
    """

    # Normalize line endings
    text = text.replace("\r", "\n")

    # Remove triple blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


# =========================================================
# DOCX EXPORT
# =========================================================

def generate_docx(answer: str, sources: list):

    doc = Document()

    # =====================================================
    # DOCUMENT TITLE
    # =====================================================

    title = doc.add_heading(
        "AI Knowledge Assistant",
        level=0
    )

    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # =====================================================
    # RESPONSE SECTION
    # =====================================================

    heading = doc.add_heading(
        "Response",
        level=1
    )

    cleaned_answer = clean_docx_markdown(answer)

    # Split into logical paragraphs
    paragraphs = cleaned_answer.split("\n\n")

    for para in paragraphs:

        para = para.strip()

        if not para:
            continue

        # -------------------------------------------------
        # BULLET POINTS
        # -------------------------------------------------

        if para.startswith("- "):

            items = para.split("\n")

            for item in items:

                item = item.replace("- ", "").strip()

                p = doc.add_paragraph(
                    item,
                    style="List Bullet"
                )

                for run in p.runs:
                    run.font.size = Pt(11)

            continue

        # -------------------------------------------------
        # NORMAL PARAGRAPH
        # -------------------------------------------------

        p = doc.add_paragraph()

        # Handle markdown bold
        parts = re.split(r"(\*\*.*?\*\*)", para)

        for part in parts:

            if part.startswith("**") and part.endswith("**"):

                run = p.add_run(
                    part[2:-2]
                )

                run.bold = True

            else:

                p.add_run(part)

        # Font size
        for run in p.runs:
            run.font.size = Pt(11)

        # Paragraph spacing
        p.paragraph_format.space_after = Pt(10)

    # =====================================================
    # SOURCES SECTION
    # =====================================================

    if sources:

        doc.add_heading(
            "Sources",
            level=1
        )

        # Create source table
        table = doc.add_table(
            rows=1,
            cols=3
        )

        table.style = "Table Grid"

        hdr_cells = table.rows[0].cells

        hdr_cells[0].text = "File"
        hdr_cells[1].text = "Chunk"
        hdr_cells[2].text = "Score"

        for src in sources:

            row_cells = table.add_row().cells

            file_name = (
                src.get("source", "Unknown")
                .split("\\")[-1]
            )

            row_cells[0].text = file_name

            row_cells[1].text = str(
                src.get("chunk_id", "-")
            )

            score = round(
                src.get("score", 0) * 100,
                1
            )

            row_cells[2].text = f"{score}%"

    # =====================================================
    # FOOTER METADATA
    # =====================================================

    doc.add_page_break()

    footer = doc.add_paragraph()

    footer_run = footer.add_run(
        "Generated by AI Knowledge Assistant"
    )

    footer_run.italic = True
    footer_run.font.size = Pt(9)

    footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # =====================================================
    # SAVE TO BUFFER
    # =====================================================

    buffer = BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return buffer